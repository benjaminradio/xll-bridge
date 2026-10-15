"""Creates a plain demo workbook whose formulas call the Lua functions in demo.lua.

Usage: python make_demo_workbook.py out.xlsx
Then embed the scripts:  xllb-pack embed out.xlsx demo.xlsx demo.lua

Formulas use the `_xll.` prefix, which is how Excel stores add-in functions in files.
"""
import sys

from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = "Demo"
ws["A1"], ws["B1"] = 2, 3
ws["A2"] = "add(A1, B1)"
ws["B2"] = '=_xll.LUA.CALL("add", A1, B1)'
ws["A3"] = "safe_div(1, 0)"
ws["B3"] = '=_xll.LUA.CALL("safe_div", 1, 0)'
ws["A4"] = "hostname()"
ws["B4"] = '=_xll.LUA.CALL("hostname")'
ws["A5"] = "eval"
ws["B5"] = '=_xll.LUA.EVAL("... * 10", A1)'
ws["A7"] = "eval, no scripts needed"
ws["B7"] = '=_xll.LUA.EVAL("1 +2")'
ws["A8"] = "hypot(3, 4) via require('lib.mathx')"
ws["B8"] = '=_xll.LUA.CALL("hypot", 3, 4)'
ws["A9"] = "slow_hostname (async, ~3 s)"
ws["B9"] = '=_xll.LUA.ASYNC("slow_hostname")'
ws["A10"] = "clock (live stream)"
ws["B10"] = '=_xll.LUA.STREAM("clock")'
ws["A6"] = "status"
ws["B6"] = "=_xll.LUA.STATUS()"
for r, (x, y) in enumerate([(1, 2), (3, 4), (5, 6)], start=11):
    ws.cell(row=r, column=1, value=x)
    ws.cell(row=r, column=2, value=y)
ws["D10"] = "transpose(A11:B13)  (Ctrl+Shift+Enter on pre-365 Excel)"
ws["D11"] = "=_xll.LUA.CALL(\"transpose\", A11:B13)"
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 18
wb.save(sys.argv[1])
